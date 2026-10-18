"""report/boat_world_model/office.py -- the index as an EXCEL workbook and the lexicon + papers as WORD documents.

    python report/boat_world_model/build_bwm.py build    # writes index/ATLAS.xlsx, LEXICON.docx, PAPERS.docx too
    python report/boat_world_model/build_bwm.py check    # ... and compares their CONTENT with a fresh build

WHY (2026-10-12, owner: "make the index in a docx and excel file (tools to make these if needed too)", and LaTeX math
in the definitions). The atlas page is for browsing; a workbook is for sorting, filtering and pasting, and a Word
document for reading, printing and annotating -- with the lexicon's equations as NATIVE Word equations, editable.

ATLAS.xlsx (openpyxl, Arial): Summary (counts as live COUNTIF formulas over the other sheets -- Excel recalculates on
open), Papers (every paper: status, held / free / library, the best link, code repo, DOI, arXiv), Lexicon (definition +
LaTeX + symbols), Citations (every verbatim quote with its paper), Labs, To download (with where-to-get-it links), Code
(the verified repos), Facts (simulators, world models, planners, RL, sensors, platform, hardware ... mined from the full
texts). Frozen header row, filters, hyperlinks.
LEXICON.docx / PAPERS.docx: pandoc from generated Markdown, so `$$...$$` becomes Office Math (OMML) -- every equation
was checked to convert (2026-10-13: 1,433 of 1,433). A table of contents by letter; terms link their related terms.
PAPERS.docx is an annotated bibliography, alphabetical by first author.

DETERMINISM: zip entries get one fixed timestamp and the document properties a fixed date (pandoc via
SOURCE_DATE_EPOCH), so a rebuild of unchanged input is byte-identical on one machine -- but zlib and pandoc versions
differ between machines, so `check` compares CONTENT (every cell value; every paragraph's text and the equation
count), never bytes. pandoc versions also write some MATH differently (2026-10-15: \\perp as U+22A5 or U+27C2 held
`check` red wherever a pandoc other than the builder's ran it), so equation text is read version-free: MATH_FOLD and
each delimiter's characters (docx_content). `build` rewrites a file only when a part changed, reading a Word file's
document.xml with those characters and 3.6.1's footnote restart folded (docx_part_canonical), so the lab's pandoc and
the laptop's no longer re-commit LEXICON.docx in turn. Without pandoc the Word files are skipped, never faked.

SOURCE STAMP (2026-10-17): every Word file carries the sha256 of the Markdown pandoc read, as the custom document
property `source_sha256` (pandoc -M writes it to docProps/custom.xml -- the same bytes from pandoc 3.1.3 and 3.6.2), so a
machine WITHOUT pandoc can still tell a Word file built from other Markdown. bed3765b rebuilt papers.json to 3,079
papers where no pandoc was; `check` said the Word files were "not compared", and origin/main stayed red on every machine
with pandoc until PAPERS.docx was rebuilt by hand. A Word file in the index must carry its fresh Markdown's stamp on
every machine (build_bwm.office_stale); one that is not there is not stale anywhere (no pandoc, no Word files).

PYTHON-DOCX (2026-10-17): the house styles (Arial body, teal headings) are python-docx's restyle of pandoc's
reference.docx (_reference_docx). Without python-docx that step fell back SILENTLY to pandoc's stock styles: the
laptop's py3.12 .venv, which lacks it, built report/ci_landing's Word files with no house styles and other closing
section properties, and `check` -- which reads text -- passed them. A machine without python-docx is now one without
pandoc for the Word files (word_missing): `build` writes none and says why, and `check` reads their stamps. python-docx
is in requirements.txt since then.
"""
from __future__ import annotations

import datetime
import hashlib
import io
import json
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path
from typing import Dict, List, Optional

FIXED = (2026, 1, 1, 0, 0, 0)
EPOCH = "1767225600"                                    # 2026-01-01T00:00:00Z, for pandoc's SOURCE_DATE_EPOCH
FONT = "Arial"
HEAD_FILL = "1F4E5F"


def _normalise_zip(data: bytes) -> bytes:
    """Every entry re-written with one fixed timestamp (same order, same compression)."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    stamp = ("%04d-%02d-%02dT%02d:%02d:%02dZ" % FIXED).encode()
    with zipfile.ZipFile(out, "w") as dst:
        for info in src.infolist():
            zi = zipfile.ZipInfo(info.filename, date_time=FIXED)
            zi.compress_type = info.compress_type
            zi.external_attr = info.external_attr
            body = src.read(info.filename)
            if info.filename == "docProps/core.xml":    # openpyxl stamps the SAVE time here, whatever was set
                body = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:)",
                              lambda m: m.group(1) + stamp + m.group(2), body)
            dst.writestr(zi, body)
    return out.getvalue()


# --------------------------------------------------------------------------------------------------- Excel
def workbook(papers: List[dict], lex: List[dict], labs: dict, facts: Dict[str, dict], code: Dict[str, dict],
             hunt: callable, meta: Optional[Dict[str, dict]] = None, *, doc: dict):
    """The openpyxl Workbook (not saved). `hunt(paper)` -> [(label, url)] (paper_links.hunt_links). `doc`: the
    project's titles and where its builder sits (build_bwm.doc_text)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
    head_fill = PatternFill("solid", fgColor=HEAD_FILL)
    body_font = Font(name=FONT, size=10)
    link_font = Font(name=FONT, size=10, color="1F5FA8", underline="single")
    wrap = Alignment(wrap_text=True, vertical="top")
    top = Alignment(vertical="top")

    def sheet(name, cols, rows, widths, links=(), wraps=()):
        ws = wb.create_sheet(name)
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill, c.alignment = head_font, head_fill, Alignment(vertical="center")
        for r in rows:
            ws.append(["" if v is None else (ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v) for v in r])
            for c in ws[ws.max_row]:
                if c.data_type == "f":                   # "=..." in a title / quote is text, never a formula
                    c.data_type = "s"
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
        link_cols = {cols.index(c) + 1 for c in links}
        wrap_cols = {cols.index(c) + 1 for c in wraps}
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = body_font
                c.alignment = wrap if c.column in wrap_cols else top
                if c.column in link_cols and isinstance(c.value, str) and c.value.startswith("http"):
                    c.hyperlink = c.value
                    c.font = link_font
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        return ws

    yes = lambda b: "yes" if b else "no"
    first = lambda p: (p.get("authors") or [""])[0]
    best = lambda p: next((u for lab, u in hunt(p) if lab not in ("Scholar", "PDF search", "Semantic Scholar")), "")
    prow = []
    for p in papers:
        c = code.get(p["key"]) or {}
        prow.append([p["key"], p["title"], first(p), "; ".join(a for a in p.get("authors") or [] if a != "et al."),
                     p["year"], p["venue"], p["status"], yes(p["published"]), p["relevance"], ", ".join(p["topics"]),
                     yes(p["pdf_held"]), p.get("file", ""), p.get("availability", ""), best(p),
                     c.get("url", "") if c.get("counted") else "", c.get("status", ""),
                     p.get("code_stars") if p.get("code_stars") is not None else c.get("stars"),
                     c.get("licence", ""), p.get("published_date", ""), p.get("cited_by"), p.get("cited_by_s2"),
                     p.get("influential_citations"), p.get("cited_here"), p.get("cites_here"),
                     p.get("venue_citedness_2yr"), p.get("venue_h_index"), p.get("meta_measured_on", ""),
                     f"https://doi.org/{p['doi']}" if p.get("doi") else "",
                     f"https://arxiv.org/abs/{p['arxiv_id']}" if p.get("arxiv_id") else "", p.get("open_url", ""),
                     p.get("summary", ""), next(iter((p.get("why_for_us") or {}).values()), "")])
    pcols = ["key", "title", "first author", "authors", "year", "venue", "status", "published", "relevance",
             "topics", "PDF held", "file", "availability", "best link", "code repo", "code status", "stars",
             "licence", "published (date)", "cited by (OpenAlex)", "cited by (Semantic Scholar)",
             "influential citations (S2)", "cited by papers in this index", "cites papers in this index",
             "venue 2-yr mean citedness", "venue h-index", "OpenAlex measured on", "DOI link", "arXiv link",
             "free copy", "summary", "why it matters here"]
    sheet("Papers", pcols, prow, [24, 60, 20, 30, 6, 28, 22, 9, 9, 28, 8, 40, 11, 40, 40, 18, 7, 12, 11, 9, 9, 9, 9,
                                  9, 9, 8, 11, 34, 34, 34, 70, 70], links=("best link", "code repo", "DOI link", "arXiv link", "free copy"),
          wraps=())
    lrow = [[x["id"], x["term"], ", ".join(x["aliases"]), ", ".join(x["categories"]), ", ".join(x["topics"]),
             x["definition"], "\n\n".join(m["latex"] for m in x.get("math") or []),
             "\n\n".join(m.get("symbols_tex") or m["symbols"] for m in x.get("math") or []), len(x["citations"]),
             ", ".join(x["related"])]
            for x in lex]
    sheet("Lexicon", ["id", "term", "aliases", "categories", "topics", "definition", "equations (LaTeX)", "symbols",
                      "citations", "related"], lrow, [26, 32, 34, 22, 26, 80, 60, 60, 9, 40],
          wraps=("definition", "equations (LaTeX)", "symbols"))
    crow = [[x["term"], c["quote"], c["key"], c["title"], c["year"], c["status"], c["where"], c["source"], c["url"]]
            for x in lex for c in x["citations"]]
    sheet("Citations", ["term", "quote (verbatim)", "paper key", "paper title", "year", "status", "where", "source",
                        "link"], crow, [30, 60, 24, 60, 6, 22, 20, 10, 40], links=("link",), wraps=("quote (verbatim)",))
    deg: Dict[str, int] = {}
    for e in labs.get("edges") or []:
        deg[e["a"]] = deg.get(e["a"], 0) + 1
        deg[e["b"]] = deg.get(e["b"], 0) + 1
    brow = [[n["lab"], n["kind"], n["pi"], n["institution"], n["country"], len(n["papers"]), deg.get(n["id"], 0),
             ", ".join(n["topics"]), n["url"]] for n in labs.get("nodes") or []]
    sheet("Labs", ["group", "kind", "PI", "institution", "country", "papers", "links", "topics", "web"], brow,
          [40, 11, 24, 40, 14, 7, 6, 40, 40], links=("web",))
    rank = {"core": 0, "high": 1, "medium": 2, "low": 3}
    todo = sorted((p for p in papers if not p["pdf_held"]),
                  key=lambda p: (rank.get(p["relevance"], 4), bool(p.get("arxiv_id")), p["key"]))
    drow = []
    for p in todo:
        links = dict((lab, u) for lab, u in reversed(hunt(p)))
        drow.append(["priority" if p["relevance"] in ("core", "high") else "later", p["key"], p["title"], p["year"],
                     p["relevance"], "on arXiv" if p.get("arxiv_id") else p.get("availability", ""), best(p),
                     links.get("publisher", ""), links.get("Scholar", ""), links.get("PDF search", "")])
    sheet("To download", ["list", "key", "title", "year", "relevance", "availability", "best link", "publisher",
                          "Scholar", "PDF search"], drow, [9, 24, 60, 6, 9, 11, 40, 34, 34, 34],
          links=("best link", "publisher", "Scholar", "PDF search"))
    by_key = {p["key"]: p for p in papers}
    stars = lambda k, c: by_key.get(k, {}).get("code_stars") if by_key.get(k, {}).get("code_stars") is not None \
        else c.get("stars")                            # one count per repo in the workbook: the dated one if measured
    krow = [[k, by_key.get(k, {}).get("title", ""), c.get("url", ""), c.get("status", ""), c.get("source", ""),
             stars(k, c), c.get("licence", ""), yes(c.get("archived")), c.get("pushed_at", "")]
            for k, c in sorted(code.items())]
    sheet("Code", ["key", "title", "repo", "status", "source", "stars", "licence", "archived", "last push"], krow,
          [24, 60, 44, 20, 8, 7, 12, 8, 11], links=("repo",))
    j = lambda f, k: ", ".join(f.get(k) or [])
    frow = [[k, by_key.get(k, {}).get("title", ""), j(f, "simulators"), j(f, "world_models"), j(f, "planners"),
             j(f, "rl"), j(f, "sensors"), j(f, "platform"), j(f, "hardware"), j(f, "datasets"), j(f, "metrics"),
             j(f, "tested"), j(f, "code_own")] for k, f in sorted(facts.items())]
    sheet("Facts", ["key", "title", "simulators", "world models", "planners", "RL", "sensors", "platform", "hardware",
                    "datasets", "metrics", "tested", "own code (stated)"], frow,
          [24, 50, 28, 28, 24, 22, 30, 22, 20, 30, 28, 18, 40])
    # `meta` (meta_links.py): which index papers cite which, and the most-cited works citing each held paper
    meta = meta or {}
    title = lambda k: by_key.get(k, {}).get("title", "")
    crow = [[k, title(k), c, title(c)] for k, m in sorted(meta.items()) for c in m.get("refs_in_corpus") or []]
    sheet("Cites in index", ["citing key", "citing title", "cited key", "cited title"], crow, [24, 60, 24, 60])
    trow = [[k, title(k), t.get("title", ""), t.get("year"), t.get("cited_by"),
             f"https://doi.org/{t['doi']}" if t.get("doi") else "", (m.get("measured") or {}).get("top_citing", "")]
            for k, m in sorted(meta.items()) for t in m.get("top_citing") or []]
    sheet("Top citing works", ["key", "title", "citing work", "year", "its citations", "DOI link", "measured on"],
          trow, [24, 50, 70, 6, 9, 34, 11], links=("DOI link",))
    # the Summary sheet: live formulas over the sheets above (Excel recalculates on open)
    ws = wb["Sheet"]
    ws.title = "Summary"
    wb.move_sheet("Summary", offset=-(len(wb.sheetnames) - 1))
    rows = [(doc["workbook"], None),
            (f"Generated by {doc['engine']}/build_bwm.py build ({doc['repo']}); every sheet has a "
             "filter row. The counts below are formulas over the other sheets.", None),
            ("", None),
            ("Papers", "=COUNTA(Papers!A:A)-1"), ("  published (peer-reviewed etc.)", '=COUNTIF(Papers!H:H,"yes")'),
            ("  preprints", '=COUNTIF(Papers!G:G,"preprint")'), ("  PDF held", '=COUNTIF(Papers!K:K,"yes")'),
            ("  with a code repo", '=COUNTIF(Papers!O:O,"http*")'),
            ("Papers to download", "=COUNTA('To download'!A:A)-1"),
            ("  priority (core / high)", "=COUNTIF('To download'!A:A,\"priority\")"),
            ("Lexicon terms", "=COUNTA(Lexicon!A:A)-1"), ("  with equations", '=COUNTIF(Lexicon!G2:G1048576,"?*")'),
            ("Verbatim citations", "=COUNTA(Citations!A:A)-1"),
            ("Research groups", '=COUNTIF(Labs!B:B,"group")'), ("Papers with mined facts", "=COUNTA(Facts!A:A)-1"),
            ("Citations between index papers", "=COUNTA('Cites in index'!A:A)-1")]
    for r in rows:
        ws.append(list(r))
    ws["A1"].font = Font(name=FONT, bold=True, size=14, color=HEAD_FILL)
    ws["A2"].font = Font(name=FONT, size=10, italic=True)
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A2:B2")
    ws.row_dimensions[2].height = 30
    for row in ws.iter_rows(min_row=4):
        for c in row:
            c.font = Font(name=FONT, size=10, bold=not str(row[0].value or "").startswith("  "))
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 14
    wb.calculation.fullCalcOnLoad = True
    when = datetime.datetime(*FIXED)
    wb.properties.created = wb.properties.modified = when
    wb.properties.creator = f"{doc['repo']} build_bwm.py"
    wb.properties.title = doc["title"]
    return wb


def xlsx_bytes(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return _normalise_zip(buf.getvalue())


def xlsx_content(wb_or_path) -> List[tuple]:
    """(sheet, row values ...) for every row of every sheet -- what `check` compares. A path or bytes is opened
    read-only and CLOSED again: a read-only workbook holds its file open until closed, and left to the garbage collector
    it kept ATLAS.xlsx open after `check` (2026-10-15: Windows would not delete or replace the file, and a late close
    under `pytest -W error` failed a run whose tests had passed). A Workbook passed in is the caller's: left open."""
    from openpyxl import load_workbook
    opened = isinstance(wb_or_path, (str, Path, bytes))
    wb = wb_or_path if not opened else load_workbook(
        io.BytesIO(wb_or_path) if isinstance(wb_or_path, bytes) else wb_or_path, read_only=True)
    try:
        out = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                out.append((ws.title,) + tuple("" if v is None else v for v in row))
        return out
    finally:
        if opened:
            wb.close()


# ---------------------------------------------------------------------------------------------------- Word
def _md(s) -> str:
    """Plain text safe in pandoc Markdown (no accidental emphasis, math, links or raw HTML)."""
    s = str(s or "")
    return re.sub(r"([\\`*_{}\[\]<>#$|~^])", lambda m: "\\" + m.group(1), s)


_INLINE = re.compile(r"(?<!\\)(\$[^$]+?(?<!\\)\$)")


def _md_math(s) -> str:
    """Like _md, but `$...$` spans (the symbols explanations' inline LaTeX) pass through as math, so pandoc writes
    them as inline Office Math; a written \\$ stays a literal dollar."""
    out = []
    for part in _INLINE.split(str(s or "")):
        if len(part) > 2 and part.startswith("$") and part.endswith("$"):
            out.append(part)
        else:
            out.append(_md(part.replace("\\$", "$")))
    return "".join(out)


def _yaml_text(s: str) -> str:
    """A metadata value for pandoc's YAML block: as it is when YAML reads it as that text (BlueBoat's titles, so its
    Word files keep their bytes), else double-quoted -- "Reef Atlas: lexicon" unquoted is a YAML mapping error."""
    s = str(s)
    if re.fullmatch(r"[A-Za-z][\w .,'()/+-]*[\w.)]", s) and \
            s.lower() not in ("true", "false", "yes", "no", "on", "off", "null"):
        return s
    return json.dumps(s)


def _front(title: str, subtitle: str) -> List[str]:
    return ["---", f"title: {_yaml_text(title)}"] + ([f"subtitle: {_yaml_text(subtitle)}"] if subtitle else []) + ["---"]


def lexicon_markdown(lex: List[dict], stats: dict, *, doc: dict) -> str:
    L = _front(doc["lexicon_docx"], doc["lexicon_docx_subtitle"]) + [
         "",
         _md(f"{stats['terms']} terms; {stats.get('math_equations', 0)} equations on {stats.get('math_terms', 0)} of "
             f"them; {stats['citations']} verbatim citations, each from an API-verified paper (published or preprint "
             f"as marked). Generated by {doc['engine']}/build_bwm.py -- edit the sources, not this file."), ""]
    letter = None
    for x in lex:
        first = x["term"][:1].upper()
        first = first if first.isalpha() else "0-9"
        if first != letter:
            letter = first
            L += [f"# {letter}", ""]
        L += [f"## {_md(x['term'])} {{#{x['id']}}}", ""]
        meta = [", ".join(x["categories"])] if x["categories"] else []
        if x["aliases"]:
            meta.append("also: " + ", ".join(x["aliases"]))
        if meta:
            L += ["*" + _md(" -- ".join(meta)) + "*", ""]
        L += [_md_math(x["definition_tex"]) if x.get("definition_tex") else _md(x["definition"]), ""]
        for m in x.get("math") or []:
            L += ["$$", m["latex"], "$$", ""]
            if m["symbols"]:
                L += ["*where* " + _md_math(m.get("symbols_tex") or m["symbols"]), ""]
        for s in x["senses"]:
            L += ["- Sense: " + _md(s)]
        others = [d for d in x["definitions"] if d["text"] != x["definition"]]
        for d in others:
            L += [f"- ({_md(d['topic'].replace('_', ' '))}) {_md(d['text'])}"]
        if x["senses"] or others:
            L.append("")
        for c in x["citations"]:
            L += [f"> \u201c{_md(c['quote'])}\u201d -- [{_md(c['key'])}]({c['url']}) "
                  f"({c['year'] or 'n.d.'}, {_md(c['status'])}; {_md(c['where'])})", ""]
        if x["related_ids"]:
            names = {y["id"]: y["term"] for y in lex}
            L += ["Related: " + ", ".join(f"[{_md(names[r])}](#{r})" for r in x["related_ids"] if r in names), ""]
    return "\n".join(L) + "\n"


def papers_markdown(papers: List[dict], stats: dict, code: Dict[str, dict], hunt, *, doc: dict) -> str:
    from unicodedata import normalize

    def sortkey(p):
        a = (p.get("authors") or [""])[0]
        sur = a.split(",")[0] if "," in a else (a.split()[-1] if a.split() else p["key"])
        return (normalize("NFKD", sur).lower(), p.get("year") or 0, p["title"].lower())
    L = _front(doc["papers_docx"], doc["papers_docx_subtitle"]) + [
         "",
         _md(f"{stats['papers']} papers ({stats['published']} published, {stats['preprints']} preprints), alphabetical "
             f"by first author; {stats.get('pdfs_held', 0)} held as PDFs. Generated by build_bwm.py."), ""]
    letter = None
    for p in sorted(papers, key=sortkey):
        s = sortkey(p)[0][:1].upper()
        s = s if s.isalpha() else "0-9"
        if s != letter:
            letter = s
            L += [f"# {letter}", ""]
        au = ", ".join(a for a in (p.get("authors") or [])[:6] if a != "et al.") + (" et al." if len(p.get("authors") or []) > 6 or "et al." in (p.get("authors") or []) else "")
        L += [f"**{_md(p['title'])}**  ", f"{_md(au) or 'n.a.'} ({p['year'] or 'n.d.'}). *{_md(p['venue'] or 'n.p.')}*. "
              f"{_md(p['status'])}; relevance {p['relevance']}; {'PDF held' if p['pdf_held'] else 'not held'}. "
              f"`{p['key']}`", ""]
        links = [(lab, u) for lab, u in hunt(p) if lab not in ("PDF search",)][:4]
        c = code.get(p["key"]) or {}
        if c.get("counted"):
            links.append((f"code ({c.get('status', '').lower().replace('_', ' ')})", c["url"]))
        if links:
            L += [" -- ".join(f"[{_md(lab)}]({u})" for lab, u in links), ""]
        if p.get("summary"):
            L += [_md(p["summary"]), ""]
    return "\n".join(L) + "\n"


def pandoc_ok() -> bool:
    return shutil.which("pandoc") is not None


def word_missing() -> Optional[str]:
    """What this machine lacks to write the Word files -- "pandoc", "python-docx" (the house styles: _reference_docx)
    -- or None when it has both. Either one missing: no Word file is written (docx_bytes), never one in pandoc's stock
    styles."""
    if not pandoc_ok():
        return "pandoc"
    try:
        import docx                                    # noqa: F401 -- python-docx (requirements.txt)
        from docx.shared import Pt, RGBColor           # noqa: F401
    except ImportError:
        return "python-docx"
    return None


def _reference_docx(tmp: Path) -> Optional[Path]:
    """pandoc's default reference.docx with the house fonts (Arial body, headings in the atlas's teal); None when pandoc
    cannot print it or python-docx cannot restyle it -- then docx_bytes writes nothing (no stock-styled Word file)."""
    ref = tmp / "reference.docx"
    r = subprocess.run(["pandoc", "-o", str(ref), "--print-default-data-file", "reference.docx"], capture_output=True)
    if r.returncode or not ref.is_file():
        return None
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
        d = Document(str(ref))
        for name, size, bold, color in (("Normal", 10, False, None), ("Body Text", 10, False, None),
                                        ("First Paragraph", 10, False, None), ("Title", 22, True, "1F4E5F"),
                                        ("Subtitle", 13, False, "52676F"), ("Heading 1", 16, True, "1F4E5F"),
                                        ("Heading 2", 12, True, "2B6A86"), ("Block Text", 9.5, False, "52676F")):
            try:
                st = d.styles[name]
            except KeyError:
                continue
            st.font.name, st.font.size, st.font.bold = FONT, Pt(size), bold
            if color:
                st.font.color.rgb = RGBColor.from_string(color)
        d.save(str(ref))
    except Exception:                                    # noqa: BLE001 -- python-docx missing or failing: no styles
        return None
    return ref


STAMP = "source_sha256"                                 # the Word files' custom property naming their Markdown


def source_stamp(markdown: str) -> str:
    """The stamp a Word file built from `markdown` carries (docx_bytes): "sha256:" and its UTF-8's hex digest."""
    return "sha256:" + hashlib.sha256(markdown.encode("utf-8")).hexdigest()


def docx_stamp(data_or_path) -> Optional[str]:
    """The source stamp a .docx carries (docProps/custom.xml, as pandoc -M writes it); None when it has none or is not
    a readable zip."""
    try:
        raw = data_or_path if isinstance(data_or_path, bytes) else Path(data_or_path).read_bytes()
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            xml = z.read("docProps/custom.xml").decode("utf-8")
    except Exception:                                    # noqa: BLE001 -- missing, not a zip, no such part: no stamp
        return None
    m = re.search(rf'<property\b[^>]*\bname="{STAMP}"[^>]*>\s*<vt:lpwstr>([^<]*)</vt:lpwstr>', xml)
    return m.group(1) if m else None


def docx_bytes(markdown: str, toc: bool = True) -> Optional[bytes]:
    """Markdown -> .docx through pandoc (math -> OMML) in the house styles, stamped with the Markdown's source_stamp;
    None when pandoc or python-docx is missing (word_missing) or either fails."""
    if word_missing():
        return None
    with tempfile.TemporaryDirectory() as td:
        tmp = Path(td)
        (tmp / "in.md").write_text(markdown, encoding="utf-8", newline="\n")
        ref = _reference_docx(tmp)
        if ref is None:
            return None
        cmd = ["pandoc", str(tmp / "in.md"), "-f", "markdown+tex_math_dollars-implicit_figures", "-t", "docx",
               "-o", str(tmp / "out.docx"), "-M", f"{STAMP}={source_stamp(markdown)}"] + \
              (["--toc", "--toc-depth=1"] if toc else []) + ["--reference-doc", str(ref)]
        r = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", errors="replace",
                           env={**os.environ, "SOURCE_DATE_EPOCH": EPOCH})
        if r.returncode or not (tmp / "out.docx").is_file():
            return None
        return _normalise_zip((tmp / "out.docx").read_bytes())


# Math characters pandoc's math library (texmath) writes differently from one version to the next, folded to one form
# before `check` compares. MEASURED 2026-10-15 over the atlas's 1,433 equations and a 683-snippet survey, between the
# committed LEXICON.docx (the lab's pandoc 3.6), the laptop's 3.6.2 and WSL Ubuntu's 3.1.3: \perp is U+22A5 UP TACK
# before texmath 0.12.8.13 and U+27C2 PERPENDICULAR from it (pandoc 3.6 -> 3.6.2), and a bare \lVert x \rVert is
# U+2225 text in 3.1.3 but a U+2016 delimiter in 3.6+. Nothing else differed. Math only: prose is compared exactly.
MATH_FOLD = str.maketrans({"\u27c2": "\u22a5", "\u2016": "\u2225"})

_PARA_TOKEN = re.compile(r"<([wm]):t(?: [^>]*)?>([^<]*)</[wm]:t>"           # a run's text: prose (w) or math (m)
                         r"|<m:(beg|sep|end)Chr\b([^>]*?)/>"                 # a delimiter's character
                         r"|<(/?)m:(d|e)(?=[\s/>])[^>]*?(/?)>")              # a delimiter object; an argument


def _paragraph_text(p: str) -> str:
    """One paragraph's text in reading order: prose as written, math through MATH_FOLD, and each delimiter object's
    characters (<m:d>: begChr, sepChr between its arguments, endChr -- OMML's defaults "(", "|", ")") where Word draws
    them. pandoc keeps those in attributes, and texmath has moved bars between text and delimiters across versions
    (\\lVert, above), so a bar must read the same either way."""
    out: List[str] = []
    open_d: List[dict] = []                            # the <m:d>s open here: their begChr's slot in `out`, sep, end
    depth = 0                                          # <m:e> elements open
    for m in _PARA_TOKEN.finditer(p):
        ns, text, which, attrs, close, tag, empty = m.groups()
        if text is not None:
            out.append(text if ns == "w" else text.translate(MATH_FOLD))
        elif which:
            if open_d:
                val = re.search(r'm:val="([^"]*)"', attrs)
                ch = (val.group(1) if val else "").translate(MATH_FOLD)
                if which == "beg":
                    out[open_d[-1]["slot"]] = ch
                else:
                    open_d[-1][which] = ch
        elif tag == "d" and not empty:
            if close:
                if open_d:
                    out.append(open_d.pop()["end"])
            else:
                open_d.append({"slot": len(out), "sep": "|", "end": ")", "depth": depth, "args": 0})
                out.append("(")
        elif tag == "e":
            if close:
                depth -= 1
                continue
            if open_d and open_d[-1]["depth"] == depth:  # an argument of the innermost delimiter, not of a part in it
                if open_d[-1]["args"]:
                    out.append(open_d[-1]["sep"])
                open_d[-1]["args"] += 1
            if not empty:
                depth += 1
    return "".join(out)


def docx_content(data_or_path) -> List[str]:
    """Every paragraph's text -- prose AND equation text (<w:t> and <m:t>, in document order, with each delimiter's
    characters; _paragraph_text) and each hyperlink's target -- plus one line counting the Office Math equations: what
    `check` compares."""
    raw = data_or_path if isinstance(data_or_path, bytes) else Path(data_or_path).read_bytes()
    z = zipfile.ZipFile(io.BytesIO(raw))
    xml = z.read("word/document.xml").decode("utf-8")
    try:
        rels = z.read("word/_rels/document.xml.rels").decode("utf-8")
    except KeyError:
        rels = ""
    target = {m.group(1): m.group(2) for m in re.finditer(r'<Relationship [^>]*?Id="([^"]+)"[^>]*?Target="([^"]*)"', rels)}
    target.update({m.group(2): m.group(1) for m in re.finditer(r'<Relationship [^>]*?Target="([^"]*)"[^>]*?Id="([^"]+)"',
                                                               rels)})
    paras = re.findall(r"<w:p[ >].*?</w:p>", xml, re.S)
    texts = []
    for p in paras:
        t = _paragraph_text(p)
        links = [target.get(r, r) for r in re.findall(r'<w:hyperlink [^>]*?r:id="([^"]+)"', p)]
        texts.append(t + "".join(f" <{u}>" for u in links))
    return [t for t in texts if t.strip()] + [f"equations: {xml.count('<m:oMath>') + xml.count('<m:oMath ')}"]


# What two pandocs write differently in one Markdown's word/document.xml, folded before `build` compares a Word file
# with the committed one (build_bwm.same_zip_payload). MEASURED 2026-10-15: the atlas's LEXICON.docx as the lab's pandoc
# 3.6 wrote it (fc0def68) and as the laptop's 3.6.2 writes it differ ONLY in that part, in six \perp glyphs (MATH_FOLD)
# and in the closing section properties -- pandoc 3.6.1 restarts footnote numbering in each section by default (#2773;
# these documents have no footnotes), where 3.6 and 3.1.3 write an empty <w:sectPr />. WSL's 3.1.3 also differs in 9 of
# the 16 parts (styles, numbering, theme, fonts, settings, their order), so its build is written whatever this folds.
_MATH_CHARS = re.compile(r'(<m:t(?: [^>]*)?>)([^<]*)(?=</m:t>)|(<m:(?:beg|sep|end)Chr\b[^>]*?m:val=")([^"]*)(?=")')
_SECT_PR = re.compile(r"<w:sectPr>(.*?)</w:sectPr>", re.S)
_FOOTNOTE_RESTART = re.compile(r'\s*<w:footnotePr>\s*<w:numRestart w:val="eachSect"\s*/>\s*</w:footnotePr>')


def docx_part_canonical(name: str, body: bytes) -> bytes:
    """A part of a pandoc-built .docx as `build` compares it: word/document.xml with only what pandoc versions write
    differently folded to one form -- each equation character (<m:t>, a delimiter's) through MATH_FOLD, 3.6.1's
    footnote restart out of the section properties -- and every other part as it is. Prose, formatting and structure
    stay byte for byte: stricter than docx_content, which `check` reads."""
    if name != "word/document.xml":
        return body

    def section(m):
        rest = _FOOTNOTE_RESTART.sub("", m.group(1))
        return f"<w:sectPr>{rest}</w:sectPr>" if rest.strip() else "<w:sectPr />"
    xml = _MATH_CHARS.sub(lambda m: m.group(1) + m.group(2).translate(MATH_FOLD) if m.group(1) is not None
                          else m.group(3) + m.group(4).translate(MATH_FOLD), body.decode("utf-8"))
    return _SECT_PR.sub(section, xml).encode("utf-8")
